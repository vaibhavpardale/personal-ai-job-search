"""Profile and document storage.

The profile is a single structured record in SQLite. Uploaded documents (for example a
résumé) are saved as files in the local documents folder, with metadata rows in SQLite.

Only facts you have provided are used as starting values. Everything else stays blank so
no details are invented.
"""

import os
import re
import uuid
from pathlib import Path

from db import PROJECT_ROOT, list_experience

DOCS_DIR_ENV_VAR = "JOB_TRACKER_DOCS_DIR"

PROFILE_FIELDS = (
    "full_name",
    "headline",
    "summary",
    "years_experience",
    "education",
    "skills",
    "achievements",
    "target_roles",
    "industries",
    "preferred_locations",
)

DOCUMENT_KINDS = ("Resume", "Other")

# Starting values taken from the background you provided. Edit them in the app.
SEED_PROFILE = {
    "full_name": "",
    "headline": "Principal Technical Product Manager at Ingram Micro",
    "summary": "",
    "years_experience": "Approximately 12 years in product management",
    "education": "IIM PGDM",
    "skills": "",
    "achievements": "",
    "target_roles": "",
    "industries": "",
    "preferred_locations": "",
}


def documents_dir():
    """Folder for uploaded documents. JOB_TRACKER_DOCS_DIR overrides it (used by tests)."""
    return Path(os.environ.get(DOCS_DIR_ENV_VAR, PROJECT_ROOT / "documents"))


def get_profile(conn):
    """Return the saved profile, or the seed values if none has been saved yet.

    The result also includes 'experience': a list of roles from your LinkedIn export and any
    entries you added. The AI uses these as facts, so they are kept separate from free text.
    """
    row = conn.execute("SELECT * FROM profile WHERE id = 1").fetchone()
    profile = dict(SEED_PROFILE) if row is None else {field: row[field] for field in PROFILE_FIELDS}
    profile["experience"] = list_experience(conn)
    return profile


def save_profile(conn, **fields):
    """Save the given profile fields. Fields you do not pass keep their current values.

    Unknown field names raise ValueError.
    """
    unknown = set(fields) - set(PROFILE_FIELDS)
    if unknown:
        raise ValueError(f"Unknown profile fields: {sorted(unknown)}")
    current = get_profile(conn)
    values = {field: fields.get(field, current[field]).strip() for field in PROFILE_FIELDS}
    columns = ", ".join(PROFILE_FIELDS)
    placeholders = ", ".join("?" for _ in PROFILE_FIELDS)
    updates = ", ".join(f"{field} = excluded.{field}" for field in PROFILE_FIELDS)
    conn.execute(
        f"""
        INSERT INTO profile (id, {columns}, updated_at)
        VALUES (1, {placeholders}, datetime('now'))
        ON CONFLICT(id) DO UPDATE SET {updates}, updated_at = datetime('now')
        """,
        [values[field] for field in PROFILE_FIELDS],
    )
    conn.commit()


def safe_filename(name):
    """Reduce an uploaded filename to a safe basename with no folders or odd characters."""
    base = Path(name).name
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._")
    return cleaned or "file"


def save_document(conn, kind, original_name, content):
    """Write an uploaded file to the documents folder and record it. Returns the row id."""
    if kind not in DOCUMENT_KINDS:
        raise ValueError(f"Unknown document kind: {kind!r}")
    folder = documents_dir()
    folder.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex[:8]}_{safe_filename(original_name)}"
    (folder / stored_name).write_bytes(content)
    cursor = conn.execute(
        "INSERT INTO documents (kind, original_name, stored_name) VALUES (?, ?, ?)",
        (kind, Path(original_name).name, stored_name),
    )
    conn.commit()
    return cursor.lastrowid


def list_documents(conn):
    """Return saved document records, newest first."""
    rows = conn.execute("SELECT * FROM documents ORDER BY uploaded_at DESC, id DESC").fetchall()
    return [dict(row) for row in rows]


def extract_text(path):
    """Return the plain text of a stored document. Supports TXT, MD, PDF, and DOCX."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document

        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    raise ValueError(f"Cannot read text from a {suffix or 'no-extension'} file")


def latest_resume(conn):
    """Return the most recently uploaded résumé record, or None."""
    return next((d for d in list_documents(conn) if d["kind"] == "Resume"), None)


def latest_resume_text(conn):
    """Return the text of the most recent résumé, or an empty string if none is uploaded."""
    resume = latest_resume(conn)
    return document_text(conn, resume["id"]) if resume else ""


def document_text(conn, doc_id):
    """Return the extracted text of a saved document by id."""
    row = conn.execute("SELECT stored_name FROM documents WHERE id = ?", (doc_id,)).fetchone()
    if row is None:
        raise ValueError(f"No document with id {doc_id}")
    return extract_text(documents_dir() / row["stored_name"])
