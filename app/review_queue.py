"""Daily review queue: search, match, tailor, and wait for your approval.

This module never submits applications, sends messages, or contacts anyone. Approving a queue
item only records your decision. Opening the posting and submitting it is always done by you.
"""

import time

from docx import Document

import ai
import db
import discovery
import profile_store

DEFAULT_MIN_SCORE = 70
DEFAULT_MAX_ANALYSES = 10


def queue_root():
    """Folder that holds the prepared files for each queued job."""
    return profile_store.documents_dir() / "queue"


def write_resume_docx(tailored, path):
    """Write the tailored résumé section as an editable Word file."""
    document = Document()
    document.add_heading(tailored["headline"], level=1)
    document.add_paragraph(tailored["summary"])
    for bullet in tailored["bullets"]:
        document.add_paragraph(bullet, style="List Bullet")
    document.save(path)


def write_review_notes(tailored, path):
    """Write what you must check before using the draft: unsupported requirements and inferences."""
    lines = ["NOT SUPPORTED BY YOUR MATERIALS (do not claim these)"]
    lines += [f"- {item}" for item in tailored["unsupported_requirements"]] or ["- None flagged"]
    lines += ["", "PLEASE CONFIRM BEFORE USING"]
    lines += [f"- {item}" for item in tailored["items_to_confirm"]] or ["- Nothing flagged"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def prepare_files(job, tailored, root):
    """Write the résumé, cover letter, and review notes for one job. Returns the folder path."""
    folder = root / f"job_{job['id']}"
    folder.mkdir(parents=True, exist_ok=True)
    write_resume_docx(tailored, folder / "tailored_resume.docx")
    (folder / "cover_letter.txt").write_text(tailored["cover_letter"] + "\n", encoding="utf-8")
    write_review_notes(tailored, folder / "review_notes.txt")
    return folder


def scan_watchlist(
    conn,
    boards,
    keywords,
    locations=(),
    exclude=(),
    fetch=discovery.fetch_board,
    pause=time.sleep,
):
    """Read the watchlist, record what is still listed, and return the postings that pass filters.

    Postings on a successfully read board that are no longer listed are marked expired.
    Returns {'jobs': [...], 'errors': [...]}.
    """
    raw, errors, read_sources, seen = discovery.collect_boards(boards, fetch=fetch, pause=pause)
    db.record_board_scan(conn, read_sources, seen)
    kept = discovery.filter_jobs(raw, keywords, locations=locations, exclude=exclude)
    return {"jobs": kept, "errors": errors}


def run_daily(
    conn,
    boards,
    keywords,
    client,
    profile,
    resume_text,
    locations=(),
    exclude=(),
    min_score=DEFAULT_MIN_SCORE,
    max_analyses=DEFAULT_MAX_ANALYSES,
    fetch=discovery.fetch_board,
    pause=time.sleep,
    root=None,
):
    """Search the watchlist, score new postings, and queue strong matches.

    At most max_analyses postings are sent to the AI per run, to keep costs predictable.
    Postings beyond that stay in the tracker and are picked up by the next run.
    Returns a summary dict for display.
    """
    root = root or queue_root()
    summary = {"found": 0, "new": 0, "duplicates": 0, "analyzed": 0, "queued": 0, "errors": []}

    scan = scan_watchlist(
        conn, boards, keywords, locations=locations, exclude=exclude, fetch=fetch, pause=pause
    )
    jobs = scan["jobs"]
    summary["found"] = len(jobs)
    summary["errors"].extend(scan["errors"])
    for job in jobs:
        outcome = db.import_discovered_job(conn, job)
        if outcome == "new":
            summary["new"] += 1
        elif outcome == "duplicate":
            summary["duplicates"] += 1

    candidates = db.jobs_needing_fit(conn)[:max_analyses]
    for job in candidates:
        try:
            analysis = ai.analyze_job_fit(client, profile, job)
            db.save_fit_analysis(conn, job["id"], analysis)
            summary["analyzed"] += 1
            if analysis["score"] < min_score:
                continue
            tailored = ai.tailor_application(client, profile, job, resume_text)
            folder = prepare_files(job, tailored, root)
            db.add_queue_item(conn, job["id"], analysis["score"], analysis["verdict"], folder)
            summary["queued"] += 1
        except Exception as exc:  # one bad posting or API error must not stop the run
            summary["errors"].append(f"{job['title']} at {job['company']}: {exc}")
    return summary
